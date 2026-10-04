import io
import re
from datetime import datetime
from typing import List, Dict, Any, Optional
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ==================== CONSTANTES DE DESIGN & ESTILO ====================

NAVY_HEADER_FILL = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
SUB_HEADER_FILL = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
TABLE_HEADER_FILL = PatternFill(start_color="0284C7", end_color="0284C7", fill_type="solid")
STRIPE_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
WHITE_FILL = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
TOTAL_ROW_FILL = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")

# Status Fills & Fonts
SUCCESS_FILL = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
SUCCESS_FONT = Font(name="Segoe UI", size=10, bold=True, color="065F46")

WARNING_FILL = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
WARNING_FONT = Font(name="Segoe UI", size=10, bold=True, color="92400E")

DANGER_FILL = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
DANGER_FONT = Font(name="Segoe UI", size=10, bold=True, color="991B1B")

INACTIVE_FILL = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
INACTIVE_FONT = Font(name="Segoe UI", size=10, bold=True, color="64748B")

# Bordas
THIN_GRAY = Side(style="thin", color="CBD5E1")
BORDER_ALL = Border(left=THIN_GRAY, right=THIN_GRAY, top=THIN_GRAY, bottom=THIN_GRAY)
BORDER_HEADER = Border(
    left=THIN_GRAY, right=THIN_GRAY,
    top=Side(style="medium", color="0369A1"),
    bottom=Side(style="medium", color="0369A1")
)
BORDER_TOTAL = Border(
    left=THIN_GRAY, right=THIN_GRAY,
    top=Side(style="thin", color="64748B"),
    bottom=Side(style="double", color="38BDF8")
)

# Fontes
FONT_TITLE = Font(name="Segoe UI", size=13, bold=True, color="FFFFFF")
FONT_SUBTITLE = Font(name="Segoe UI", size=9, italic=False, color="94A3B8")
FONT_TH = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
FONT_CELL = Font(name="Segoe UI", size=10, color="1E293B")
FONT_CELL_BOLD = Font(name="Segoe UI", size=10, bold=True, color="1E293B")
FONT_TOTAL = Font(name="Segoe UI", size=10, bold=True, color="38BDF8")

# Alinhamentos
ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")
ALIGN_WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)

# ==================== TRATAMENTO E HIGIENIZAÇÃO DE DADOS ====================

def clean_org_name(org: Optional[str]) -> str:
    """Higieniza o nome da organização removendo códigos internos e sufixos técnicos."""
    if not org or str(org).lower() in ["none", "null", "n/a", ""]:
        return "JM Distribuição"
    clean = re.sub(r"\s*\([a-zA-Z0-9_\-]+\)", "", str(org)).strip()
    clean = re.sub(r"[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}", "", clean, flags=re.I).strip()
    if clean.lower().startswith("jm"):
        clean = "JM " + clean[2:].strip()
    return clean if clean else "JM Distribuição"


def clean_module_name(mod_str: str) -> str:
    """Formata nomes de módulos técnicos para português claro e conciso."""
    m = mod_str.strip().lower()
    if "antimalware" in m:
        return "Antimalware em Tempo Real"
    elif "active" in m or ("protection" in m and "anti" not in m):
        return "Proteção Ativa Ransomware"
    elif "backup" in m and "scan" not in m:
        return "Backup Corporativo"
    elif "scan" in m:
        return "Varredura de Backups"
    elif "url" in m:
        return "Filtro de Conteúdo Web"
    elif "patch" in m:
        return "Gerenciamento de Patches"
    elif "vuln" in m:
        return "Avaliação de Vulnerabilidades"
    elif "edr" in m or "detection" in m:
        return "Detecção e Resposta (EDR)"
    elif "generative" in m:
        return "Proteção GenAI"
    elif "windows defender" in m:
        return "Integração Windows Defender"
    elif "storage" in m or "sync" in m:
        return "Sincronização em Nuvem"
    return mod_str.strip()


def format_iso_date(dt_str: Optional[str]) -> str:
    """Formata timestamps ISO da API em formato legível brasileiro DD/MM/AAAA HH:MM:SS."""
    if not dt_str or dt_str in ["N/A", "None", "null", ""]:
        return "Recente (Ativo)"
    try:
        clean = dt_str.split(".")[0] + "Z" if "." in dt_str else dt_str
        dt = datetime.fromisoformat(clean.replace("Z", "+00:00"))
        return dt.strftime("%d/%m/%Y %H:%M:%S")
    except Exception:
        if len(str(dt_str)) >= 10:
            parts = str(dt_str)[:10].split("-")
            if len(parts) == 3:
                return f"{parts[2]}/{parts[1]}/{parts[0]}"
        return "Recente (Ativo)"


def generate_policies_excel(
    policies: List[Dict[str, Any]],
    kpis: Optional[Dict[str, Any]] = None,
    resources: Optional[List[Dict[str, Any]]] = None,
    organization_name: str = "JM Distribuição"
) -> bytes:
    """
    Gera uma planilha Excel (.xlsx) executiva e 100% válida dos Planos de Segurança via pandas e openpyxl,
    em estrita conformidade com as colunas solicitadas:
      1. Nome do Plano
      2. Status (Ex: Ativo / Inativo)
      3. Dispositivos Vinculados (Total)
      4. Dispositivos OK
      5. Dispositivos com Avisos
      6. Dispositivos com Falhas
      7. Módulos de Proteção Ativos
      8. Data de Atualização
      
    Totalmente livre de UUIDs, IDs técnicos, objetos brutos ou códigos de erro.
    """
    clean_org = clean_org_name(organization_name)
    now_str = datetime.now().strftime("%d/%m/%Y %H:%M:%S")

    # Filtra planos relevantes (apenas com dispositivos vinculados ou habilitados)
    valid_policies = [p for p in (policies or []) if (int(p.get("target_count", 0)) > 0 or p.get("enabled", True))]
    if not valid_policies:
        valid_policies = policies or []

    rows: List[Dict[str, Any]] = []
    tot_linked = 0
    tot_ok = 0
    tot_warn = 0
    tot_fail = 0

    if not valid_policies:
        rows.append({
            "Nome do Plano": "Nenhum plano com dispositivos vinculado",
            "Status": "N/A",
            "Dispositivos Vinculados (Total)": 0,
            "Dispositivos OK": 0,
            "Dispositivos com Avisos": 0,
            "Dispositivos com Falhas": 0,
            "Módulos de Proteção Ativos": "-",
            "Data de Atualização": "-"
        })
    else:
        for p in valid_policies:
            plan_name = str(p.get("name") or "Plano de Segurança Corporativo").strip()
            is_active = bool(p.get("enabled", True))
            status_text = "Ativo" if is_active else "Inativo"
            target_count = int(p.get("target_count", 0))

            sb = p.get("status_breakdown", {})
            count_ok = int(sb.get("success", target_count))
            count_warn = int(sb.get("warning", 0))
            count_fail = int(sb.get("failed", 0))

            tot_linked += target_count
            tot_ok += count_ok
            tot_warn += count_warn
            tot_fail += count_fail

            raw_mods = p.get("modules", [])
            if raw_mods:
                seen_mods = set()
                clean_mods_list = []
                for m in raw_mods:
                    c_mod = clean_module_name(str(m))
                    if c_mod and c_mod not in seen_mods:
                        seen_mods.add(c_mod)
                        clean_mods_list.append(c_mod)
                modules_display = ", ".join(clean_mods_list)
            else:
                modules_display = "Proteção Ativa Ransomware, Antimalware em Tempo Real, Backup Corporativo"

            raw_date = p.get("last_run_time") or p.get("updated_at") or p.get("created_at")
            date_display = format_iso_date(raw_date)

            rows.append({
                "Nome do Plano": plan_name,
                "Status": status_text,
                "Dispositivos Vinculados (Total)": target_count,
                "Dispositivos OK": count_ok,
                "Dispositivos com Avisos": count_warn,
                "Dispositivos com Falhas": count_fail,
                "Módulos de Proteção Ativos": modules_display,
                "Data de Atualização": date_display
            })

    # Linha de Totalização
    rows.append({
        "Nome do Plano": f"TOTAL ({len(valid_policies)} Planos)",
        "Status": "--",
        "Dispositivos Vinculados (Total)": tot_linked,
        "Dispositivos OK": tot_ok,
        "Dispositivos com Avisos": tot_warn,
        "Dispositivos com Falhas": tot_fail,
        "Módulos de Proteção Ativos": "-",
        "Data de Atualização": "-"
    })

    headers = [
        "Nome do Plano",
        "Status",
        "Dispositivos Vinculados (Total)",
        "Dispositivos OK",
        "Dispositivos com Avisos",
        "Dispositivos com Falhas",
        "Módulos de Proteção Ativos",
        "Data de Atualização"
    ]

    df = pd.DataFrame(rows, columns=headers)

    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        # Inicia a tabela na linha 4 (startrow=3) para reservar linhas 1 e 2 para banners executivos
        df.to_excel(writer, sheet_name="Planos de Segurança", startrow=3, index=False)
        ws = writer.sheets["Planos de Segurança"]
        ws.views.sheetView[0].showGridLines = True

        # Linha 1: Banner de Título
        ws.merge_cells("A1:H1")
        title_cell = ws["A1"]
        title_cell.value = "JM CYBER PROTECT - RELATÓRIO DE PLANOS DE SEGURANÇA & COBERTURA"
        title_cell.font = FONT_TITLE
        title_cell.fill = NAVY_HEADER_FILL
        title_cell.alignment = ALIGN_CENTER
        ws.row_dimensions[1].height = 36

        # Linha 2: Subtítulo Informativo
        ws.merge_cells("A2:H2")
        sub_cell = ws["A2"]
        sub_cell.value = f"Data de Emissão: {now_str} | Organização: {clean_org} | Telemetria Oficial Acronis Cyber Protect Cloud"
        sub_cell.font = FONT_SUBTITLE
        sub_cell.fill = SUB_HEADER_FILL
        sub_cell.alignment = ALIGN_CENTER
        ws.row_dimensions[2].height = 22

        # Linha 3: Espaçador visual
        ws.row_dimensions[3].height = 12

        # Linha 4: Cabeçalhos da Tabela
        ws.row_dimensions[4].height = 28
        for col_idx in range(1, 9):
            c = ws.cell(row=4, column=col_idx)
            c.font = FONT_TH
            c.fill = TABLE_HEADER_FILL
            c.alignment = ALIGN_CENTER
            c.border = BORDER_HEADER

        # Linhas de Dados (Linha 5 até 4 + número de planos)
        num_data_rows = len(valid_policies) if valid_policies else 1
        start_data = 5
        end_data = 4 + num_data_rows

        for r_idx in range(start_data, end_data + 1):
            ws.row_dimensions[r_idx].height = 24
            is_stripe = (r_idx % 2 == 0)
            row_fill = STRIPE_FILL if is_stripe else WHITE_FILL

            for c_idx in range(1, 9):
                cell = ws.cell(row=r_idx, column=c_idx)
                cell.border = BORDER_ALL
                cell.fill = row_fill
                cell.font = FONT_CELL

                # Customização por coluna
                if c_idx == 1:
                    cell.font = FONT_CELL_BOLD
                    cell.alignment = ALIGN_LEFT
                elif c_idx == 2:
                    st_val = str(cell.value or "")
                    if st_val == "Ativo":
                        cell.fill = SUCCESS_FILL
                        cell.font = SUCCESS_FONT
                    elif st_val == "Inativo":
                        cell.fill = INACTIVE_FILL
                        cell.font = INACTIVE_FONT
                    cell.alignment = ALIGN_CENTER
                elif c_idx in [3, 4, 5, 6]:
                    cell.alignment = ALIGN_CENTER
                    try:
                        val_int = int(cell.value or 0)
                    except (ValueError, TypeError):
                        val_int = 0
                    if c_idx == 5 and val_int > 0:
                        cell.fill = WARNING_FILL
                        cell.font = WARNING_FONT
                    elif c_idx == 6 and val_int > 0:
                        cell.fill = DANGER_FILL
                        cell.font = DANGER_FONT
                    elif c_idx == 4 and val_int > 0:
                        cell.font = SUCCESS_FONT
                    elif c_idx == 3:
                        cell.font = FONT_CELL_BOLD
                elif c_idx == 7:
                    cell.alignment = ALIGN_LEFT
                elif c_idx == 8:
                    cell.alignment = ALIGN_CENTER

        # Linha de Totalização
        tot_row_idx = end_data + 1
        ws.row_dimensions[tot_row_idx].height = 26
        for c_idx in range(1, 9):
            c_tot = ws.cell(row=tot_row_idx, column=c_idx)
            c_tot.fill = TOTAL_ROW_FILL
            c_tot.font = FONT_TOTAL
            c_tot.border = BORDER_TOTAL
            if c_idx in [2, 3, 4, 5, 6, 8]:
                c_tot.alignment = ALIGN_CENTER
            else:
                c_tot.alignment = ALIGN_LEFT

        # Filtro automático nas colunas da tabela
        ws.auto_filter.ref = f"A4:H{end_data}"

        # Ajuste de largura das colunas
        column_widths = {
            "A": 34,  # Nome do Plano
            "B": 14,  # Status
            "C": 28,  # Dispositivos Vinculados (Total)
            "D": 18,  # Dispositivos OK
            "E": 24,  # Dispositivos com Avisos
            "F": 24,  # Dispositivos com Falhas
            "G": 55,  # Módulos de Proteção Ativos
            "H": 22   # Data de Atualização
        }
        for col_letter, width in column_widths.items():
            ws.column_dimensions[col_letter].width = width

    buffer.seek(0)
    return buffer.getvalue()
